from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path(__file__).parent/'HUONG_DAN_SU_DUNG_A4.docx'
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.12
for name,size,color in [('Title',19,RGBColor(11,45,93)),('Heading 1',14,RGBColor(15,66,141)),('Heading 2',12,RGBColor(22,89,157))]:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=color
 st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(5)
p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('HƯỚNG DẪN SỬ DỤNG VÀ TRIỂN KHAI\nEDU EXAM A4')
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Phòng thi trực tuyến • Nhập Word • Duyệt đề • Chấm điểm').italic=True

def h(t):d.add_heading(t,1)
def p(t):d.add_paragraph(t)
def item(t):d.add_paragraph(t,style='List Bullet')

h('I. Hai chế độ sử dụng')
p('Chế độ GIÁO VIÊN nhập và quản lý đề thi, duyệt đáp án, mở phòng, chấm bài và xuất Excel. Chế độ HỌC SINH làm bài và chỉ xem thông báo xác nhận sau khi nộp. Chức năng trực tuyến cần có Supabase và hosting HTTPS.')
h('II. Thao tác khi chưa có máy chủ')
for t in ['Mở tệp index.html (bản độc lập) trong Chrome hoặc Edge.', 'Chọn GIÁO VIÊN → MỞ BỘ NHẬP WORD CHƯA KẾT NỐI.', 'Chọn tệp đề thi có đuôi .docx. Không tải PDF hoặc .doc vào nút Word.', 'Duyệt lần lượt tất cả câu, đối chiếu màu đỏ, công thức, hình và đáp án với tệp gốc.', 'Nếu thiếu công thức/hình WMF, chèn ảnh gốc PNG/JPEG đúng nội dung rồi xử lý cảnh báo.', 'Nhập điểm câu và quy tắc chấm có căn cứ; xuất JSON sao lưu khi chưa kết nối máy chủ.']:item(t)
h('III. Chuẩn đề Word theo mẫu đã thống nhất')
for t in ['PHẦN I – TRẮC NGHIỆM NHIỀU LỰA CHỌN: các lựa chọn A/B/C/D; chữ đỏ chỉ phương án đúng.', 'PHẦN II – TRẮC NGHIỆM ĐÚNG/SAI: từng ý a/b/c/d; màu đỏ là chỉ báo ban đầu cần xác nhận; chọn quy tắc điểm theo hướng dẫn chấm.', 'PHẦN III – TRẮC NGHIỆM TRẢ LỜI NGẮN: dòng “Đáp án:” chứa giá trị cần nhập chính xác.', 'PHẦN IV – TỰ LUẬN (nếu có): cần bổ sung rubric/biểu điểm và giáo viên chấm.']:item(t)
p('Bộ đề thực tế đang dùng để kiểm thử có 15 câu PHẦN I, 2 câu PHẦN II và 4 câu PHẦN III. Các hình WMF trong đề được báo thiếu, không được tự thay bằng hình minh họa. Điểm câu không được tự chia khi thiếu hướng dẫn chấm.')
h('IV. Bật chế độ online bằng Supabase')
for t in ['Tạo Supabase Project tại supabase.com/dashboard.', 'Cài Supabase CLI, mở thư mục PHONG_THI_A4 chứa thư mục supabase.', 'Chạy lệnh: supabase login; supabase link --project-ref YOUR_PROJECT_REF; supabase db push; supabase functions deploy exam-api.', 'Trong Supabase Auth tạo giáo viên, sau đó cấp role “teacher” cho người dùng trong app_profiles (theo README.md).', 'Đăng tải toàn bộ thư mục web lên hosting tĩnh có HTTPS.', 'Tại WebApp, mở KẾT NỐI, điền Project URL và publishable key; bấm KIỂM TRA MÁY CHỦ.', 'Tuyệt đối không đưa service_role/secret key vào HTML, tài liệu hoặc khóa công khai.']:item(t)
h('V. Quy trình giáo viên khi đã kết nối')
for t in ['ĐĂNG NHẬP bằng tài khoản giáo viên đã được cấp quyền.', 'NHẬP WORD/JSON → kiểm tra từng câu và các cảnh báo, nhập biểu điểm.', 'Duyệt từng câu: chỉ đánh dấu xác nhận nếu nội dung, hình, đáp án, quy tắc chấm đã chính xác.', 'LƯU NGÂN HÀNG ĐỀ: đề được lưu trên máy chủ; tệp Word gốc được gửi vào Storage riêng tư khi nhập từ Word.', 'TẠO PHÒNG: chọn đề, nhập thời gian làm bài và khoảng thời gian phòng mở.', 'Chia sẻ mã phòng, link hoặc mã QR qua môi trường HTTPS.', 'XEM BẢNG ĐIỂM → CHẤM TỰ LUẬN → XUẤT EXCEL. Điểm tạm tính và điểm chính thức được phân biệt.']:item(t)
h('VI. Quy trình học sinh')
for t in ['Mở link phòng thi hoặc trang HỌC SINH và nhập mã phòng.', 'Nhập họ tên và lớp trước khi bắt đầu.', 'Làm bài và chuyển câu bằng nút điều hướng; đánh dấu câu cần xem lại nếu cần.', 'Hệ thống tự lưu câu trả lời lên máy chủ; bản lưu tạm trong trình duyệt hỗ trợ khôi phục.', 'Nhấn NỘP BÀI hoặc hệ thống kích hoạt khi hết giờ. Sau nộp chỉ hiển thị thông báo thành công.']:item(t)
h('VII. Kiểm tra và giới hạn cần biết')
for t in ['Word chứa WMF/EMF hoặc công thức OMML phức tạp phải bổ sung/đối chiếu thủ công.', 'Hệ thống chỉ chấm theo đáp án và biểu điểm giáo viên xác nhận.', 'Bản online chưa được triển khai khi chưa có dự án Supabase và tên miền do người sử dụng cấu hình.', 'Cần kiểm thử đồng thời nhiều thiết bị, kiểm tra sai lệch thời gian, mạng yếu, phân quyền và kế hoạch sao lưu trước kỳ thi thật.', 'Đăng nhập học sinh bằng họ tên và lớp chưa chống được việc giả mạo; triển khai thực tế nên bổ sung định danh và giới hạn tốc độ truy cập.']:item(t)
p('Mã nguồn và các bước cấu hình chuyên sâu xem tệp README.md trong gói PHONG_THI_A4.')
for paragraph in d.paragraphs:
 paragraph.paragraph_format.widow_control=True
 if paragraph.style.name.startswith('Heading'):paragraph.paragraph_format.keep_with_next=True
 if paragraph.style.name=='List Bullet':paragraph.paragraph_format.left_indent=Cm(0.65)
d.save(out)
print(out, out.stat().st_size)
