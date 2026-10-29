"""Single-node educational video app. Deploy behind HTTPS/reverse proxy and auth."""
import json
import os
import re
import tempfile
from pathlib import Path
from contextlib import asynccontextmanager
from typing import Optional
from fastapi import FastAPI,HTTPException,Request,Depends,UploadFile,File
from fastapi.responses import FileResponse,JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel,Field
import settings as cfg
import store
import workflow as wf

TOKEN_SCHEME=HTTPBearer(auto_error=False)

def auth(credentials: Optional[HTTPAuthorizationCredentials]=Depends(TOKEN_SCHEME)):
    import hmac
    if cfg.APP_ACCESS_TOKEN and (not credentials or not hmac.compare_digest(credentials.credentials,cfg.APP_ACCESS_TOKEN)):
        raise HTTPException(401,'Cần đăng nhập bằng mã truy cập ứng dụng (không phải API key).')
    return True

@asynccontextmanager
async def lifespan(app):
    store.init()
    # Previously interrupted tasks remain persisted; do not blindly resubmit billable jobs.
    yield

app=FastAPI(title='EDUVIDEO FLOW AI V4.1.1',lifespan=lifespan)
ROOT=Path(__file__).resolve().parent

@app.exception_handler(wf.WorkflowError)
async def wf_error(request,exc):return JSONResponse({'detail':str(exc)},status_code=409)

@app.get('/',include_in_schema=False)
@app.get('/index.html',include_in_schema=False)
def index():return FileResponse(ROOT/'index.html',media_type='text/html; charset=utf-8')

@app.get('/api/health')
def health():
    import shutil
    return {'ok':True,'ai_ready':bool(cfg.GEMINI_API_KEY),'veo_project_configured':bool(cfg.GOOGLE_CLOUD_PROJECT),'access_token_protected':bool(cfg.APP_ACCESS_TOKEN),'ffmpeg_configured':bool(shutil.which(cfg.FFMPEG) and shutil.which(cfg.FFPROBE))}

@app.get('/api/projects',dependencies=[Depends(auth)])
def list_projects():return store.all_projects()

class ProjectInput(BaseModel):
    subject:str=Field(min_length=1,max_length=100)
    grade:int=Field(ge=1,le=12)
    title:str=Field(min_length=1,max_length=180)
    content:str=Field(min_length=1)
    video_type:str=Field(min_length=1)
    duration:int
    style:str=Field(min_length=1)
    character_mode:str=Field(min_length=1)
    subtitles:bool=False

@app.post('/api/projects',dependencies=[Depends(auth)])
def create_project(body:ProjectInput):
    project=wf.new_project(body.model_dump())
    wf.enqueue(project['id'],'full')
    return wf.summarize(store.get(project['id']))

@app.get('/api/projects/{project_id}',dependencies=[Depends(auth)])
def get_project(project_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    return wf.summarize(p)

@app.put('/api/projects/{project_id}',dependencies=[Depends(auth)])
def update_project(project_id:str,body:ProjectInput):
    return wf.summarize(wf.update_input(project_id,body.model_dump()))

@app.post('/api/projects/{project_id}/resume',dependencies=[Depends(auth)])
def resume(project_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    if p['status']=='COMPLETED':raise HTTPException(409,'Dự án đã hoàn thành')
    return wf.summarize(wf.enqueue(project_id,'full'))

class SceneEdit(BaseModel):
    patch:dict[str,str]

@app.patch('/api/projects/{project_id}/scenes/{scene_id}',dependencies=[Depends(auth)])
def edit_scene(project_id:str,scene_id:str,body:SceneEdit):
    return wf.summarize(wf.edit_scene(project_id,scene_id,body.patch))

@app.post('/api/projects/{project_id}/scenes/{scene_id}/regenerate-prompts',dependencies=[Depends(auth)])
def regenerate_prompts(project_id:str,scene_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    wf.find_scene(p,scene_id)
    if not p.get('bible'):raise HTTPException(409,'Chưa tạo và khóa Bible')
    return wf.summarize(wf.enqueue(project_id,'prompts',scene_id))

@app.post('/api/projects/{project_id}/scenes/{scene_id}/accept-image',dependencies=[Depends(auth)])
def accept_image(project_id:str,scene_id:str):
    return wf.summarize(wf.accept_scene_image(project_id,scene_id))

@app.post('/api/projects/{project_id}/scenes/{scene_id}/regenerate-image',dependencies=[Depends(auth)])
def regenerate_image(project_id:str,scene_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    if project_id in wf._active:raise HTTPException(409,'Dự án đang xử lý')
    s=wf.find_scene(p,scene_id)
    if not s.get('image_prompt') or not s.get('video_prompt') or not p.get('bible'):raise HTTPException(409,'Chưa đủ prompt ảnh, prompt video và Bible')
    s['accepted']=False;s['video_file']=None;s['render_job']=None;s['render_fingerprint']=None;s['status']='IMAGE_READY' if s.get('image_file') else 'DRAFT'
    p['video']=None;p['steps']['13']=p['steps']['14']='LOCKED';p['version']+=1
    wf.refresh_scene_steps(p);store.save(p)
    return wf.summarize(wf.enqueue(project_id,'image',scene_id))

@app.post('/api/projects/{project_id}/scenes/{scene_id}/regenerate-video',dependencies=[Depends(auth)])
def regenerate_video(project_id:str,scene_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    if project_id in wf._active:raise HTTPException(409,'Dự án đang xử lý')
    s=wf.find_scene(p,scene_id)
    if not s.get('image_file') or not s.get('video_prompt'):raise HTTPException(409,'Cần ảnh và prompt video hợp lệ')
    s['video_file']=None;s['accepted']=False;s['status']='IMAGE_READY';s['render_job']=None;s['render_fingerprint']=None
    p['video']=None;p['steps']['13']=p['steps']['14']='LOCKED';p['version']+=1
    wf.refresh_scene_steps(p);store.save(p)
    return wf.summarize(wf.enqueue(project_id,'video',scene_id))

@app.post('/api/projects/{project_id}/scenes/{scene_id}/accept',dependencies=[Depends(auth)])
def accept_scene(project_id:str,scene_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    with wf._guard:
        if project_id in wf._active:raise HTTPException(409,'Dự án đang xử lý')
        s=wf.find_scene(p,scene_id)
        if s['status'] not in ['VIDEO_READY','ACCEPTED'] or not s.get('video_file') or not Path(s['video_file']).is_file():raise HTTPException(409,'Cảnh chưa có video đạt kiểm tra')
        try:
            from media import probe
            probe(s['video_file'])
        except Exception as e:raise HTTPException(409,str(e)) from e
        s['accepted']=True;s['status']='ACCEPTED';store.save(p)
    return wf.summarize(p)

@app.post('/api/projects/{project_id}/merge',dependencies=[Depends(auth)])
def merge_project(project_id:str):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    if not p['scenes'] or not all(s['accepted'] and s['status']=='ACCEPTED' and s.get('video_file') and Path(s['video_file']).is_file() for s in p['scenes']):
        raise HTTPException(409,'Chỉ ghép khi tất cả cảnh đã có video hợp lệ và được chấp nhận')
    return wf.summarize(wf.enqueue(project_id,'merge'))

def serve_media(project_id,scene_id,kind):
    p=store.get(project_id)
    if not p:raise HTTPException(404,'Không có dự án')
    s=wf.find_scene(p,scene_id)
    f=s.get(kind+'_file')
    if not f or not Path(f).is_file():raise HTTPException(404,'Không có tệp hợp lệ')
    base=store.file_dir(project_id).resolve()
    if not Path(f).resolve().is_relative_to(base):raise HTTPException(403,'Đường dẫn không hợp lệ')
    return FileResponse(f,media_type='image/png' if kind=='image' else 'video/mp4',headers={'Cache-Control':'no-store'})

@app.get('/api/projects/{project_id}/scenes/{scene_id}/image',dependencies=[Depends(auth)])
def scene_image(project_id:str,scene_id:str):return serve_media(project_id,scene_id,'image')

@app.get('/api/projects/{project_id}/scenes/{scene_id}/video',dependencies=[Depends(auth)])
def scene_video(project_id:str,scene_id:str):return serve_media(project_id,scene_id,'video')

@app.get('/api/projects/{project_id}/video',dependencies=[Depends(auth)])
def final_video(project_id:str):
    p=store.get(project_id)
    if not p or not p.get('video') or not Path(p['video']['file']).is_file():raise HTTPException(404,'Video cuối chưa có')
    return FileResponse(p['video']['file'],media_type='video/mp4',headers={'Cache-Control':'no-store'})

@app.get('/api/projects/{project_id}/download',dependencies=[Depends(auth)])
def download(project_id:str):
    p=store.get(project_id)
    if not p or not p.get('video') or not Path(p['video']['file']).is_file():raise HTTPException(404,'Chưa xuất được MP4')
    return FileResponse(p['video']['file'],media_type='video/mp4',filename=f'EDUVIDEO_{project_id}.mp4')

@app.get('/api/projects/{project_id}/subtitles',dependencies=[Depends(auth)])
def download_subtitles(project_id:str):
    p=store.get(project_id)
    subtitle=p.get('video',{}).get('subtitles_srt') if p and p.get('video') else None
    if not subtitle or not Path(subtitle).is_file():raise HTTPException(404,'Dự án không có phụ đề đã xuất')
    return FileResponse(subtitle,media_type='text/plain; charset=utf-8',filename=f'EDUVIDEO_{project_id}.srt')

@app.post('/api/extract',dependencies=[Depends(auth)])
async def extract_document(file:UploadFile=File(...)):
    filename=Path(file.filename or '').name
    ext=Path(filename).suffix.lower()
    if ext not in ['.txt','.docx','.pdf']:raise HTTPException(400,'Chỉ hỗ trợ TXT, DOCX, PDF')
    blob=await file.read(cfg.MAX_SOURCE_CHARS*12+1)
    if len(blob)>cfg.MAX_SOURCE_CHARS*12:raise HTTPException(413,'Tài liệu quá lớn')
    try:
        if ext=='.txt': text=blob.decode('utf-8-sig')
        elif ext=='.docx':
            from docx import Document
            from io import BytesIO
            doc=Document(BytesIO(blob))
            lines=[p.text for p in doc.paragraphs]
            for t in doc.tables:
                for r in t.rows:lines.append(' | '.join(c.text for c in r.cells))
            text='\n'.join(lines)
        else:
            from pypdf import PdfReader
            from io import BytesIO
            reader=PdfReader(BytesIO(blob))
            text='\n'.join(page.extract_text() or '' for page in reader.pages)
    except ImportError:
        raise HTTPException(503,'Server chưa cài thư viện đọc tài liệu')
    except Exception as e:raise HTTPException(422,'Tệp không đọc được: '+str(e)[:160]) from e
    text=text.strip()
    if not text:raise HTTPException(422,'Không trích xuất được chữ; PDF quét ảnh cần OCR riêng')
    if len(text)>cfg.MAX_SOURCE_CHARS:raise HTTPException(413,'Nội dung trích xuất vượt giới hạn')
    return {'text':text,'length':len(text),'filename':filename}

app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')

if __name__=='__main__':
    import uvicorn
    uvicorn.run('server:app',host='0.0.0.0',port=cfg.PORT,reload=False)
