from pathlib import Path
from functools import partial
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from threading import Thread
from playwright.sync_api import sync_playwright
from zipfile import ZipFile
from docx import Document
p=Path(__file__).resolve().parents[1]
out=Path('/mnt/data/CHEM_MAP_TUTOR_AI_2_1_Jigsaw_docs')
server=ThreadingHTTPServer(('127.0.0.1',0),partial(SimpleHTTPRequestHandler,directory=str(p)))
Thread(target=server.serve_forever,daemon=True).start()
with sync_playwright() as pw:
 browser=pw.chromium.launch(headless=True,executable_path='/usr/bin/chromium',args=['--no-sandbox','--disable-dev-shm-usage'])
 page=browser.new_page(viewport={'width':1440,'height':900},accept_downloads=True)
 errors=[]
 page.on('pageerror',lambda e: errors.append(str(e)))
 page.on('dialog',lambda d:d.accept())
 page.set_content((p/'CHEM_MAP_TUTOR_AI_2_1_Offline.html').read_text(encoding='utf8'),wait_until='domcontentloaded')
 page.wait_for_timeout(500)
 assert 'CHEM-MAP TUTOR AI' in page.title()
 page.locator('#role-switch').select_option('teacher')
 page.locator('#nav-main [data-nav="jigsaw"]').click()
 assert page.locator('#jigsaw-form').count()==1
 assert '45 phút' in page.locator('.page-title').text_content()
 assert page.locator('tbody tr').count()>=6
 page.screenshot(path=str(out/'jigsaw_desktop.png'),full_page=True)
 page.locator('[name="title"]').fill('Nitric acid – Tính acid và tính oxi hóa')
 page.locator('[name="reviewed"]').check()
 page.locator('#jigsaw-form button[type="submit"]').click()
 for mode in ['teacher','student']:
  sel='j-export-'+mode
  with page.expect_download(timeout=20000) as d:
   page.locator(f'[data-action="{sel}"]').click()
  downloaded=d.value
  dest=out/('Jigsaw_Hoa_11_'+('Giao_vien' if mode=='teacher' else 'Hoc_sinh')+'_Le_Ba_Nhan.docx')
  downloaded.save_as(dest)
  assert dest.stat().st_size>5000, dest.stat().st_size
  with ZipFile(dest) as z:
   xml=z.read('word/document.xml').decode('utf8')
   assert ('<w:document') in xml
  doc=Document(dest)
  text='\n'.join(x.text for x in doc.paragraphs)
  if mode=='student':
   assert '0,20 mol' not in text
   assert 'Hướng dẫn giải:' not in text
   assert 'KẾ HOẠCH BÀI DẠY' not in text
  else:
   assert 'KẾ HOẠCH BÀI DẠY' in text
   assert '0,20 mol' in text
  print(mode,dest.stat().st_size,'bytes',len(doc.paragraphs),'paragraphs',len(doc.tables),'tables')
 # Import an example DOCX content to source only and validate no auto-answers
 page.locator('#j-source-file').set_input_files(str(out/'Jigsaw_Hoa_11_Giao_vien_Le_Ba_Nhan.docx'))
 page.wait_for_timeout(250)
 assert page.locator('[name="source"]').input_value().startswith('[Tài liệu:')
 assert page.locator('[name="questionA-0-q"]').input_value().startswith('A1.')
 page.screenshot(path=str(out/'jigsaw_uploaded.png'),full_page=True)
 assert not errors, errors
 print('TEST PASS: 6 groups + docx teacher/student + source import + no answer leak + no page errors')
 mobile=browser.new_page(viewport={'width':390,'height':844},accept_downloads=True)
 mobile.set_content((p/'CHEM_MAP_TUTOR_AI_2_1_Offline.html').read_text(encoding='utf8'),wait_until='domcontentloaded');mobile.locator('#role-switch').select_option('teacher')
 mobile.locator('.mobile-menu').click();mobile.locator('#nav-main [data-nav="jigsaw"]').click()
 mobile.screenshot(path=str(out/'jigsaw_mobile.png'),full_page=True)
 assert mobile.locator('#jigsaw-form').count()==1
 browser.close()
server.shutdown()
