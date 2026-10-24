from pathlib import Path
from docx import Document
from docx.shared import Inches,Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image
root=Path(__file__).parent
out=Path('/mnt/data/CHEM_MAP_TUTOR_AI_2_0_Huong_Dan_Trien_Khai_Le_Ba_Nhan.docx')
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.85);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2.35);sec.right_margin=Cm(2.15)
styles=doc.styles
for key in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=styles[key];st.font.name='Times New Roman';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
n=styles['Normal'];n.font.size=Pt(11);n.font.color.rgb=RGBColor(33,48,71);n.paragraph_format.space_after=Pt(4);n.paragraph_format.line_spacing=1.14
for i,size,c in [(1,14,RGBColor(27,65,112)),(2,11.5,RGBColor(20,131,124))]:
 st=styles[f'Heading {i}'];st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=c;st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
styles['Title'].font.size=Pt(22);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=RGBColor(27,65,112)
styles['Subtitle'].font.size=Pt(11)
def shade(cell, fill):
 tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);tcPr.append(sh)
def margins(cell,top=70,start=110,bottom=70,end=110):
 tc=cell._tc;tcPr=tc.get_or_add_tcPr();x=tcPr.first_child_found_in('w:tcMar')
 if x is None:x=OxmlElement('w:tcMar');tcPr.append(x)
 for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  node=OxmlElement('w:'+side);node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa');x.append(node)
def p(text='',style=None,bold=False):
 par=doc.add_paragraph(style=style)
 r=par.add_run(text);r.bold=bold
 return par
def heading(t,level=1):doc.add_heading(t,level)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,h in enumerate(headers):
  cell=t.rows[0].cells[i];cell.text=h;shade(cell,'244875');margins(cell)
  for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True;run.font.size=Pt(9.5)
  cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for ridx,row in enumerate(rows):
  cells=t.add_row().cells
  for j,val in enumerate(row):
   cell=cells[j];cell.text=str(val);margins(cell);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ridx%2==0:shade(cell,'F5F8FC')
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(1)
    for run in par.runs:run.font.size=Pt(9.4)
 t.style='Table Grid'
 if widths:
  for row in t.rows:
   for cell,w in zip(row.cells,widths):cell.width=Cm(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
def bullet(txt):p('• '+txt)
# title
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=par.add_run('HỒ SƠ TRIỂN KHAI WEBAPP');r.bold=True;r.font.size=Pt(12);r.font.color.rgb=RGBColor(19,142,132)
par=doc.add_paragraph(style='Title');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.add_run('CHEM-MAP TUTOR AI 2.0')
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=par.add_run('HỆ SINH THÁI CHẨN ĐOÁN VÀ KHẮC PHỤC\nLỖI SAI HÓA HỌC CÁ NHÂN HÓA');r.bold=True;r.font.size=Pt(13);r.font.color.rgb=RGBColor(42,88,132)
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=par.add_run('Giải pháp ứng dụng cho Hóa học 10 – 11 – 12\nChương trình GDPT 2018');r.font.size=Pt(11)
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=par.add_run('Tác giả: Lê Bá Nhẫn – Trường THPT Hòn Đất');r.bold=True;r.font.size=Pt(12)
cap=Image.open(root/'tests/home.png');cap.crop((250,65,1410,815)).save(root/'tests/cover_crop.png')
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.add_run().add_picture(str(root/'tests/cover_crop.png'),width=Cm(15.9))
par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=par.add_run('BẢN THỰC THI 2.0 · WEBAPP CÓ THỂ CHẠY THỬ · NGÀY 24/10/2026');r.bold=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor(83,107,134)
doc.add_page_break()
heading('1. MỤC TIÊU VÀ PHẠM VI THỰC THI')
p('WebApp hỗ trợ giáo viên và học sinh hoàn thành chu trình: làm bài → chấm theo quy tắc → giải thích phương án nhiễu → đề xuất bài luyện → đo lại → lưu minh chứng. Điểm độc đáo là mã hóa lỗi sai và phân biệt bằng chứng quan sát với giả thuyết nguyên nhân.')
table(['Nhóm tính năng','Đã hoạt động trong bản 2.0','Giới hạn hiện tại'],[
('Ngân hàng học liệu','51 câu mẫu: 36 A–D, 6 Đúng/Sai, 6 trả lời ngắn, 3 tự luận.','Bộ mẫu phải kiểm duyệt trước khi sử dụng chính thức.'),
('Chấm và phản hồi','Chấm A–D, từng ý Đúng/Sai, đáp số; hiển thị điểm /10, lời giải, câu đã chọn.','Tự luận chờ giáo viên duyệt; chưa chấm AI tự do.'),
('Bản đồ lỗi sai','16 mã lỗi K01–T04; gắn phương án nhiễu với cơ chế nhầm lẫn có căn cứ.','Mã lỗi một đáp án chỉ là giả thuyết, không phải chẩn đoán đã xác thực.'),
('Học bù & minh chứng','Bài luyện đề xuất, ghi nhận bài trước/sau, thống kê kỹ năng và CSV/JSON.','Số lượng bài còn ít thì đánh giá năng lực phải thận trọng.'),
('Quản lý giáo viên','Soạn và duyệt câu hỏi, giao đề, duyệt điểm tự luận, báo cáo lớp demo.','Vai trò chưa có đăng nhập an toàn hay CSDL chung.')],[3.1,8,5.35])
heading('2. KIẾN TRÚC HOẠT ĐỘNG')
p('Đây là ứng dụng web tĩnh local-first gồm giao diện HTML/CSS/JavaScript; bộ dữ liệu câu hỏi; động cơ chấm – chẩn đoán; dữ liệu trạng thái LocalStorage. Không cần nhập GEMINI_API_KEY và không gọi dịch vụ AI từ bên ngoài ở phiên bản chạy thử.')
table(['Thành phần','Nhiệm vụ'],[
('assets/bank.js / bank.json','Chương theo khối 10–12, 51 câu mẫu, đáp án, lời giải, mã lỗi, metadata.'),
('assets/engine.js','Chấm bốn dạng câu, điểm /10, giả thuyết lỗi, ước lượng kỹ năng, đề xuất bài học bù.'),
('assets/app.js','Hiển thị giao diện theo vai trò, quản lý bài làm, xuất nhập minh chứng, rubric.'),
('assets/style.css','Giao diện responsive trên máy tính và điện thoại.'),
('index.html / Offline.html','Bản triển khai web tĩnh và bản HTML đơn tệp chạy trực tiếp.')],[4.7,11.75])
heading('3. QUY TRÌNH SỬ DỤNG DÀNH CHO HỌC SINH')
for s in [
'Bước 1. Mở file CHEM_MAP_TUTOR_AI_2_0_Offline.html bằng Chrome hoặc Edge.',
'Bước 2. Chọn vai trò “Học sinh”, sau đó chọn hồ sơ demo A/B/C (không dùng tên thật).',
'Bước 3. Vào “Luyện tập thông minh”: chọn khối, chương, hình thức, mức độ, số câu.',
'Bước 4. Làm bài và nhấn “Nộp bài & xem phân tích”.',
'Bước 5. Xem đáp án đúng, lựa chọn của mình, lời giải; chú ý nhãn “Giả thuyết từ phương án nhiễu”.',
'Bước 6. Vào “Bản đồ lỗi sai” và “Lộ trình cá nhân”, làm gói bài học bù rồi đối chiếu kết quả.'
]:bullet(s)
heading('4. QUY TRÌNH SỬ DỤNG DÀNH CHO GIÁO VIÊN')
for s in [
'Chuyển vai trò “Giáo viên” → “Không gian giáo viên” để xem lượt bài theo hồ sơ demo.',
'Vào thẻ “Soạn câu hỏi” → nhập dữ kiện, bốn phương án hoặc rubric phù hợp → “Lưu bản nháp”.',
'Kiểm tra tính chính xác khoa học; bấm “Xác nhận kiểm duyệt” để đưa câu vào kho luyện tập.',
'Vào “Đề kiểm tra” → nhập tên đề, chọn khối và bộ lọc → “Tạo và trộn đề kiểm tra”.',
'Khi có bài tự luận, vào thẻ “Duyệt tự luận” → đối chiếu bài làm/rubric → cho điểm câu từ 0 đến 1.',
'Vào “Thống kê & minh chứng” để xuất CSV; trong “Dữ liệu & hệ thống” có thể sao lưu toàn bộ JSON.'
]:bullet(s)
heading('5. CẤU TRÚC MỘT PHẢN HỒI LỖI SAI')
p('Ví dụ Hóa học 11: Cu tác dụng với HNO₃ loãng. Theo giả thiết NO là sản phẩm khử duy nhất: 3Cu + 8HNO₃(loãng) → 3Cu(NO₃)₂ + 2NO + 4H₂O.')
table(['Thành tố','Nội dung sư phạm'],[
('Căn cứ hóa học','Cu⁰ → Cu²⁺ + 2e; N(+5) trong nitrate nhận electron để thành NO (N+2).'),
('Đáp án chuẩn','Khí tạo thành theo phương trình chuẩn là NO, không phải NO₂.'),
('Lỗi nghi vấn','Chọn NO₂ có thể do nhầm HNO₃ loãng với HNO₃ đặc; mã D01.'),
('Mức tin cậy','Chỉ là giả thuyết từ lựa chọn. Cần bài giải từng bước hoặc câu hỏi kiểm tra bổ sung.'),
('Bài học bù','Nhận biết sản phẩm khử theo dữ kiện, cân bằng electron, tỉ lệ n(Cu) : n(NO) = 3 : 2.')],[4,12.45])
heading('6. NGUYÊN TẮC CHẤM ĐIỂM VÀ KIỂM ĐỊNH')
p('Trắc nghiệm A–D đúng được 1 điểm câu; Đúng/Sai đúng k/4 ý được k/4 điểm câu; trả lời ngắn đúng (theo sai số số học hợp lệ) được 1 điểm câu. Tự luận không tự động chấm, mặc định đang chờ giáo viên duyệt. Điểm bài làm = 10 × (tổng điểm câu đã đạt)/(tổng số câu); khi còn tự luận chưa duyệt là điểm tạm tính, câu chờ chấm tạm tính 0.')
p('Bản chạy thử không sử dụng phân loại lỗi ở mức “đã xác nhận” nếu học sinh chỉ chọn một phương án. Nội dung mô tả bẫy thông tin được giáo viên thiết kế trước, không phải lời khẳng định AI đã đọc được tư duy bên trong học sinh.')
heading('7. TRIỂN KHAI TRÊN MÁY CÁ NHÂN HOẶC HOSTING')
table(['Cách chạy','Hướng dẫn thao tác'],[
('Đơn giản nhất','Mở file HTML đơn tệp trong trình duyệt; toàn bộ logic, CSS và ngân hàng câu hỏi đã được đóng gói.'),
('Chạy bằng máy chủ tĩnh','Giải nén ZIP → mở terminal ở thư mục dự án → chạy python -m http.server 8080 → mở http://localhost:8080.'),
('Đưa lên Vercel/Netlify','Tạo dự án website tĩnh từ thư mục có index.html và assets/. Không cần build command; dùng thư mục gốc làm đầu ra. Tệp vercel.json đã kèm sẵn.'),
('Kiểm tra bàn giao','Chạy node tests/engine.test.js để kiểm tra động cơ; các bài kiểm tra trình duyệt có trong tests/browser_test.py và browser_extended.py.')],[4.1,12.35])
doc.add_page_break()
heading('8. KIỂM THỬ ĐÃ THỰC HIỆN')
table(['Mã thử','Tình huống','Trạng thái'],[
('T01','Chạy động cơ với 51 câu và bốn dạng câu.','Đạt'),('T02','Học sinh làm → nộp → xem điểm /10 và lời giải.','Đạt'),('T03','Giáo viên thêm câu → phê duyệt → tạo đề kiểm tra.','Đạt'),('T04','Luyện Đúng/Sai; ghi nhận gói học bù và cặp trước/sau.','Đạt'),('T05','Tự luận chờ duyệt → giáo viên chấm → cập nhật điểm.','Đạt'),('T06','Màn hình máy tính/điện thoại và điều hướng.','Đạt')],[2.3,11.45,2.7])
p('Các trạng thái kiểm thử trên xác nhận luồng chức năng trong môi trường test trình duyệt; không phải bằng chứng về hiệu quả sư phạm, độ an toàn triển khai diện rộng hoặc kiểm định độc lập tất cả nội dung Hóa học.')
heading('9. QUẢN TRỊ RỦI RO VÀ HƯỚNG NÂNG CẤP')
for s in [
'Không nhập dữ liệu định danh, điểm số thật hoặc hồ sơ riêng tư của học sinh: bản web tĩnh chỉ dùng LocalStorage và không có xác thực phân quyền.',
'Đưa vào sử dụng trong nhà trường: xây dựng backend với đăng nhập, quyền truy cập theo lớp, lưu dữ liệu an toàn, log, chính sách quyền riêng tư và phương án sao lưu.',
'Nếu tích hợp AI sinh nội dung: dùng khóa API ở phía máy chủ nếu nhà cung cấp yêu cầu; không đặt bí mật trên client; chặn xuất bản học liệu chưa qua duyệt.',
'Triển khai nghiên cứu sư phạm: đo đầu vào/đầu ra bằng bài tương đương, lưu lịch sử can thiệp, có thể dùng nhóm so sánh; không dùng số liệu demo làm kết quả.',
'Mở rộng dần ngân hàng câu hỏi theo bài/chương SGK Kết nối tri thức 10–12; kiểm duyệt công thức, điều kiện phản ứng, rubric, phương án nhiễu.'
]:bullet(s)
heading('10. GÓI BÀN GIAO')
p('01 — CHEM_MAP_TUTOR_AI_2_0_Offline.html: file WebApp một tệp, mở để trải nghiệm nhanh.')
p('02 — CHEM_MAP_TUTOR_AI_2_0_WebApp.zip: mã nguồn, ngân hàng câu hỏi, kiểm thử, README và thiết lập deploy tĩnh.')
p('03 — Tài liệu Word này: quy trình vận hành và định hướng đưa ứng dụng vào thực tế.')
p('Biên soạn và phát triển: Lê Bá Nhẫn – Trường THPT Hòn Đất. Bản trình diễn được chuẩn bị ngày 24/10/2026.')
# Footer
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('CHEM-MAP TUTOR AI 2.0 · Lê Bá Nhẫn – THPT Hòn Đất · Tài liệu triển khai WebApp');r.font.name='Times New Roman';r.font.size=Pt(8);r.font.color.rgb=RGBColor(116,130,153)
doc.save(out)
print(out,out.stat().st_size)
