from pathlib import Path
from playwright.sync_api import sync_playwright
from bs4 import BeautifulSoup
from openpyxl import load_workbook
P=Path(__file__).resolve().parent.parent
soup=BeautifulSoup((P/'web/index.html').read_text(),'html.parser')
mock=soup.new_tag('script');mock.string="var __kv={};Object.defineProperty(window,'localStorage',{configurable:true,value:{getItem:k=>__kv[k]??null,setItem:(k,v)=>{__kv[k]=v},removeItem:k=>delete __kv[k]}})";soup.head.append(mock)
for scr in list(soup.find_all('script')):
 if scr.has_attr('src'):
  filepath=P/'web'/scr['src']
  if 'mathjax' in str(filepath):scr.decompose();continue
  scr.string=filepath.read_text();del scr['src']
with sync_playwright() as p:
 b=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox','--disable-dev-shm-usage','--disable-gpu'])
 page=b.new_page(accept_downloads=True)
 errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
 page.set_content(str(soup))
 page.evaluate('''() => {const q=[{id:'111',position:1,section:'I',type:'mcq',body:'Câu kiểm thử thật trong bộ tự kiểm',points:1,options:[{id:'A',text:'Phương án A'},{id:'B',text:'Phương án B'}],statements:[],media:[]}];window.API.call=async(action,data)=>{if(action==='student.room')return {exam:{title:'Kiểm thử',subject:'Hóa học'},room:{duration_minutes:45,ends_at:new Date(Date.now()+3600000).toISOString()}};if(action==='student.join')return {token:'a'.repeat(48),attempt_id:'1',deadline_at:new Date(Date.now()+600000).toISOString(),server_time:new Date().toISOString(),questions:q};if(action==='student.save')return {ok:true};if(action==='student.submit')return {ok:true,submitted:true};return {ok:true}};}''')
 page.locator('[data-go="student"]').click()
 page.locator('#studentCode').fill('SAMPLETST')
 page.locator('#studentName').fill('Nguyễn Văn A')
 page.locator('#studentClass').fill('10A1')
 page.locator('#joinRoom').click()
 page.locator('input[type=radio][value="A"]').check()
 assert page.evaluate("window.A4_DEBUG.getState().answers['111']")=='A'
 page.on('dialog',lambda dialog:dialog.accept())
 page.locator('#submitExam').click()
 page.wait_for_function("document.querySelector('#app').textContent.includes('ĐÃ NỘP BÀI THÀNH CÔNG!')")
 assert 'Điểm chính thức' not in page.locator('#app').inner_text()
 with page.expect_download() as dl:
  page.evaluate('''async()=>{const b=await window.A4_DEBUG.gradeXMLRows([['Họ tên','Điểm'],['HS A',9.5]]);const u=URL.createObjectURL(b),a=document.createElement('a');a.href=u;a.download='test.xlsx';a.click();}''')
 file=P/'tests'/'test_export.xlsx';dl.value.save_as(file)
 ws=load_workbook(file,read_only=True).active
 assert ws['A2'].value=='HS A' and ws['B2'].value==9.5
 print('OK MOCK STUDENT: join, answer, submit without grades, valid XLSX')
 print('PAGE_ERRORS:',errors)
 assert not errors
 b.close()
